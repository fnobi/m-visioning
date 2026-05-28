from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1: Title
slide_layout = prs.slide_layouts[6]  # blank
slide1 = prs.slides.add_slide(slide_layout)

# Title
title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "m-visioning"
p.font.size = Pt(60)
p.font.bold = True
p.font.color.rgb = RGBColor(0x55, 0x77, 0xCC)
p.alignment = PP_ALIGN.CENTER

# Subtitle
sub_box = slide1.shapes.add_textbox(Inches(0.5), Inches(4), Inches(12.333), Inches(1))
tf = sub_box.text_frame
p = tf.paragraphs[0]
p.text = "個人向け資産シミュレーションWebアプリ"
p.font.size = Pt(32)
p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
p.alignment = PP_ALIGN.CENTER

# Slide 2: Features
slide2 = prs.slides.add_slide(slide_layout)

# Header
header_box = slide2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1))
tf = header_box.text_frame
p = tf.paragraphs[0]
p.text = "主な機能"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = RGBColor(0x55, 0x77, 0xCC)

# Feature 1
f1_box = slide2.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(6), Inches(2.5))
tf = f1_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "銀行口座残高予測"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

p2 = tf.add_paragraph()
p2.text = "• 複数の銀行口座を管理"
p2.font.size = Pt(18)
p2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p3 = tf.add_paragraph()
p3.text = "• 入出金履歴を記録"
p3.font.size = Pt(18)
p3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p4 = tf.add_paragraph()
p4.text = "• 将来の残高をシミュレーション"
p4.font.size = Pt(18)
p4.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p5 = tf.add_paragraph()
p5.text = "• グラフ表示で推移を可視化"
p5.font.size = Pt(18)
p5.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# Feature 2
f2_box = slide2.shapes.add_textbox(Inches(6.5), Inches(1.5), Inches(6), Inches(2.5))
tf = f2_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "クレジットカード支払い予測"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

p2 = tf.add_paragraph()
p2.text = "• 複数カードの利用額を管理"
p2.font.size = Pt(18)
p2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p3 = tf.add_paragraph()
p3.text = "• 締め日・引落日を設定"
p3.font.size = Pt(18)
p3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p4 = tf.add_paragraph()
p4.text = "• 口座からの引落を自動計算"
p4.font.size = Pt(18)
p4.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# Feature 3
f3_box = slide2.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(6), Inches(2.5))
tf = f3_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "入出金予定の管理"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

p2 = tf.add_paragraph()
p2.text = "• 定期的な収入・支出を登録"
p2.font.size = Pt(18)
p2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p3 = tf.add_paragraph()
p3.text = "• 月次/年次の繰り返し設定"
p3.font.size = Pt(18)
p3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p4 = tf.add_paragraph()
p4.text = "• カテゴリ分けで整理"
p4.font.size = Pt(18)
p4.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# Tech Stack
tech_box = slide2.shapes.add_textbox(Inches(6.5), Inches(4.2), Inches(6), Inches(2.5))
tf = tech_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "技術スタック"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

p2 = tf.add_paragraph()
p2.text = "• Next.js + React + TypeScript"
p2.font.size = Pt(18)
p2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p3 = tf.add_paragraph()
p3.text = "• Firebase (認証・データベース)"
p3.font.size = Pt(18)
p3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

p4 = tf.add_paragraph()
p4.text = "• Emotion (CSS-in-JS)"
p4.font.size = Pt(18)
p4.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# Save
prs.save('/home/user/m-visioning/m-visioning-overview.pptx')
print("Created: m-visioning-overview.pptx")
